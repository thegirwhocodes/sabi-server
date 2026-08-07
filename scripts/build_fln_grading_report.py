#!/usr/bin/env python3
"""Build the cited FLN grading research report as a polished DOCX.

Design system: documents skill `standard_business_brief` preset with the
`editorial_cover` header template.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "FLN_GRADING_RESEARCH_AND_SABI_STANDARD_2026-08-07.md"
OUTPUT = ROOT / "docs" / "FLN_GRADING_RESEARCH_AND_SABI_STANDARD_2026-08-07.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "203748"
MUTED = "64748B"
GOLD = "B68A2A"
LIGHT = "F2F4F7"
WHITE = "FFFFFF"
TABLE_DXA = 9360
TABLE_INDENT_DXA = 120


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run_font(run, *, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]):
    if sum(widths_dxa) != TABLE_DXA:
        raise ValueError(f"Table widths must total {TABLE_DXA}: {widths_dxa}")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_field(paragraph, instruction: str):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])
    return run


def create_numbering_instance(doc: Document) -> int:
    """Create a fresh real decimal list so each Markdown list restarts at 1."""
    numbering = doc.part.numbering_part.element
    abstract_id = None
    for abstract in numbering.findall(qn("w:abstractNum")):
        for style in abstract.iter(qn("w:pStyle")):
            if style.get(qn("w:val")) == "ListNumber":
                abstract_id = int(abstract.get(qn("w:abstractNumId")))
                break
        if abstract_id is not None:
            break
    if abstract_id is None:
        abstract_id = 0

    existing = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(existing, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)
    return num_id


def set_paragraph_numbering(paragraph, num_id: int):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = num_pr.find(qn("w:ilvl"))
    if ilvl is None:
        ilvl = OxmlElement("w:ilvl")
        num_pr.append(ilvl)
    ilvl.set(qn("w:val"), "0")
    num_id_el = num_pr.find(qn("w:numId"))
    if num_id_el is None:
        num_id_el = OxmlElement("w:numId")
        num_pr.append(num_id_el)
    num_id_el.set(qn("w:val"), str(num_id))


def add_hyperlink(paragraph, label: str, url: str, *, size=None, bold=False):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    if bold:
        bold_el = OxmlElement("w:b")
        r_pr.append(bold_el)
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size * 2)))
        r_pr.append(sz)
    run.append(r_pr)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(
    r"(\[[^\]]+\]\(https?://[^)]+\)|\*\*[^*]+\*\*|https?://[^\s]+|\[[0-9]+\])"
)


def add_inline(paragraph, text: str, *, size=None, color=None, bold=False, italic=False):
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, size=size, color=color, bold=bold, italic=italic)
        token = match.group(0)
        if token.startswith("[") and "](" in token:
            label, url = token[1:].split("](", 1)
            add_hyperlink(paragraph, label, url[:-1], size=size, bold=bold)
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, color=color, bold=True, italic=italic)
        elif token.startswith("http"):
            trailing = ""
            url = token
            while url and url[-1] in ".,;":
                trailing = url[-1] + trailing
                url = url[:-1]
            add_hyperlink(paragraph, url, url, size=size, bold=bold)
            if trailing:
                run = paragraph.add_run(trailing)
                set_run_font(run, size=size, color=color, bold=bold, italic=italic)
        else:
            run = paragraph.add_run(token)
            set_run_font(run, size=size, color=BLUE, bold=False, italic=italic)
            run.font.superscript = True
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=size, color=color, bold=bold, italic=italic)


def paragraph_border_bottom(paragraph, color=GOLD, size=12, space=6):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    normal.paragraph_format.widow_control = True

    tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in tokens.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.widow_control = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def configure_page(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.34)
    section.footer_distance = Inches(0.34)


def configure_cover_section(section):
    configure_page(section)
    section.different_first_page_header_footer = False
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    section.header.paragraphs[0].clear()
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Education for Equality  •  Research completed 7 August 2026")
    set_run_font(run, size=8, color=MUTED)


def configure_content_section(section):
    configure_page(section)
    section.different_first_page_header_footer = False
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False

    p = section.header.paragraphs[0]
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("EDUCATION FOR EQUALITY  |  SABI RESEARCH")
    set_run_font(run, size=8, color=MUTED, bold=True)
    paragraph_border_bottom(p, color="D7DCE2", size=6, space=4)

    p = section.footer.paragraphs[0]
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("FLN grading research  •  ")
    set_run_font(run, size=8, color=MUTED)
    page_run = add_field(p, "PAGE")
    set_run_font(page_run, size=8, color=MUTED)


def add_cover(doc: Document):
    for _ in range(3):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(24)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run("EVIDENCE REVIEW")
    set_run_font(run, size=10, color=GOLD, bold=True)
    run.font.all_caps = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run("How Leading Organizations\nAssess Foundational Learning")
    set_run_font(run, size=30, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(26)
    run = p.add_run("What their grading practices mean for Sabi’s\nfoundational literacy and numeracy architecture")
    set_run_font(run, size=15, color=BLUE)
    paragraph_border_bottom(p, color=GOLD, size=12, space=14)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(24)
    run = p.add_run("Pratham / ASER  •  TaRL  •  J-PAL  •  UNICEF  •  USAID / RTI  •\nWorld Bank  •  UNESCO  •  Nigeria NALABE  •  Save the Children  •  Room to Read")
    set_run_font(run, size=10, color=MUTED)

    table = doc.add_table(rows=3, cols=2)
    set_table_geometry(table, [2700, 6660])
    values = [
        ("Prepared for", "Education for Equality / Sabi"),
        ("Research date", "7 August 2026"),
        ("Document status", "Research synthesis and pilot recommendation"),
    ]
    for row, (label, value) in zip(table.rows, values):
        keep_row_together(row)
        shade_cell(row.cells[0], LIGHT)
        for cell in row.cells:
            cell.text = ""
        add_inline(row.cells[0].paragraphs[0], label, size=9, color=DARK_BLUE, bold=True)
        add_inline(row.cells[1].paragraphs[0], value, size=9, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(22)
    run = p.add_run("A practical assessment architecture—not a literature review alone")
    set_run_font(run, size=10, color=GOLD, bold=True, italic=True)

def table_widths(column_count: int) -> list[int]:
    maps = {
        2: [2736, 6624],
        3: [2016, 3672, 3672],
        4: [1656, 2304, 2664, 2736],
        5: [1656, 1800, 2088, 1800, 2016],
    }
    if column_count in maps:
        return maps[column_count]
    base = TABLE_DXA // column_count
    values = [base] * column_count
    values[-1] += TABLE_DXA - sum(values)
    return values


def add_markdown_table(doc: Document, lines: list[str]):
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        rows.append(cells)
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in rows[1]):
        rows.pop(1)
    if not rows:
        return

    col_count = len(rows[0])
    widths = table_widths(col_count)
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for r_idx, values in enumerate(rows):
        row = table.rows[r_idx]
        keep_row_together(row)
        if r_idx == 0:
            set_repeat_table_header(row)
        for c_idx, value in enumerate(values):
            cell = row.cells[c_idx]
            cell.text = ""
            if r_idx == 0:
                shade_cell(cell, LIGHT)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            add_inline(
                p,
                value,
                size=8.4 if col_count >= 4 else 9,
                color=DARK_BLUE if r_idx == 0 else NAVY,
                bold=(r_idx == 0),
            )
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(0)


def add_source_paragraph(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.0
    add_inline(p, text, size=9, color=NAVY)


def render_markdown(doc: Document, markdown: str):
    lines = markdown.splitlines()
    start = next(i for i, line in enumerate(lines) if line.startswith("## Executive conclusion"))
    i = start
    in_sources = False
    active_numbering_id = None
    while i < len(lines):
        raw = lines[i].rstrip()
        stripped = raw.strip()
        if not stripped:
            active_numbering_id = None
            i += 1
            continue

        if stripped.startswith("|"):
            active_numbering_id = None
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            add_markdown_table(doc, table_lines)
            continue

        heading_match = re.match(r"^(#{2,4})\s+(.*)$", stripped)
        if heading_match:
            active_numbering_id = None
            hashes, title = heading_match.groups()
            level = {2: 1, 3: 2, 4: 3}[len(hashes)]
            p = doc.add_paragraph(style=f"Heading {level}")
            add_inline(p, title, color=BLUE if level < 3 else DARK_BLUE, bold=True)
            if title == "Executive conclusion":
                p.paragraph_format.space_before = Pt(0)
            in_sources = title == "Sources"
            i += 1
            continue

        bullet_match = re.match(r"^-\s+(.*)$", stripped)
        number_match = re.match(r"^\d+\.\s+(.*)$", stripped)
        if bullet_match or number_match:
            text = (bullet_match or number_match).group(1)
            p = doc.add_paragraph(style="List Bullet" if bullet_match else "List Number")
            if number_match:
                if active_numbering_id is None:
                    active_numbering_id = create_numbering_instance(doc)
                set_paragraph_numbering(p, active_numbering_id)
            else:
                active_numbering_id = None
            add_inline(p, text, color=NAVY)
            i += 1
            continue

        if in_sources and re.match(r"^\[\d+\]", stripped):
            active_numbering_id = None
            add_source_paragraph(doc, stripped)
            i += 1
            continue

        paragraph_lines = [stripped]
        i += 1
        while i < len(lines):
            candidate = lines[i].strip()
            if not candidate:
                break
            if candidate.startswith(("##", "|", "- ")) or re.match(r"^\d+\.\s+", candidate):
                break
            paragraph_lines.append(candidate)
            i += 1
        text = " ".join(paragraph_lines)
        active_numbering_id = None
        p = doc.add_paragraph()
        if in_sources and text.startswith("All web sources"):
            p.paragraph_format.space_before = Pt(8)
            add_inline(p, text, size=9, color=MUTED, italic=True)
        else:
            add_inline(p, text, color=NAVY)


def add_document_metadata(doc: Document):
    props = doc.core_properties
    props.title = "How Leading Organizations Assess Foundational Learning—and What Sabi Should Do"
    props.subject = "Foundational literacy and numeracy assessment research and proposed Sabi grading standard"
    props.author = "Education for Equality"
    props.keywords = "Sabi, FLN, assessment, grading, mastery, TaRL, ASER, EGRA, EGMA, UNICEF, numeracy, literacy"
    props.comments = "Generated from a cited Markdown source using the standard_business_brief design preset."


def main():
    doc = Document()
    doc.settings.odd_and_even_pages_header_footer = False
    configure_styles(doc)
    configure_cover_section(doc.sections[0])
    add_document_metadata(doc)
    add_cover(doc)
    content_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_content_section(content_section)
    render_markdown(doc, SOURCE.read_text(encoding="utf-8"))

    final_p = doc.add_paragraph()
    final_p.paragraph_format.space_before = Pt(12)
    paragraph_border_bottom(final_p, color=GOLD, size=8, space=4)
    final_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = final_p.add_run("EDUCATION FOR EQUALITY  •  SABI")
    set_run_font(run, size=8, color=MUTED, bold=True)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
